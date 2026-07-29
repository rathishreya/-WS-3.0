from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ORANGE="EA7B2C"; NAVY="112949"; SKY="3CC3F2"; SOFT="FDEEE2"; SKYSOFT="E4F6FD"
GREY="F2F4F7"; LINE="D9DEE6"; GOODS="E2EEE7"; WARNS="FEF5DC"; CRITS="F7E4E3"
F="Arial"
hdr=Font(name=F,size=9,bold=True,color="FFFFFF")
sub=Font(name=F,size=8.5,bold=True,color=NAVY)
body=Font(name=F,size=9,color="1A1A1A")
bodyb=Font(name=F,size=9,bold=True,color=NAVY)
mut=Font(name=F,size=8.5,color="5A6472")
thin=Side(style="thin",color=LINE)
bd=Border(left=thin,right=thin,top=thin,bottom=thin)
wrap=Alignment(wrap_text=True,vertical="top")
ctr=Alignment(horizontal="center",vertical="center",wrap_text=True)
wb=Workbook()

def band(ws,ncols,text,color=ORANGE,h=26,size=14):
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=ncols)
    c=ws.cell(1,1,text);c.font=Font(name=F,size=size,bold=True,color="FFFFFF")
    c.fill=PatternFill("solid",fgColor=color);c.alignment=Alignment(vertical="center",indent=1)
    ws.row_dimensions[1].height=h

# ================= START HERE =================
ws=wb.active; ws.title="START HERE"
band(ws,3,"WS 3.0 · CLIENT WORKSPACE — what must exist, and why",size=15)
ws.column_dimensions['A'].width=4;ws.column_dimensions['B'].width=30;ws.column_dimensions['C'].width=104
r=3
def blk(title,lines):
    global r
    ws.cell(r,2,title).font=Font(name=F,size=11,bold=True,color=NAVY);r+=1
    for a,b in lines:
        ws.cell(r,2,a).font=bodyb;ws.cell(r,2).fill=PatternFill("solid",fgColor=SOFT)
        ws.cell(r,2).border=bd;ws.cell(r,2).alignment=Alignment(wrap_text=True,vertical="center")
        ws.cell(r,3,b).font=body;ws.cell(r,3).border=bd;ws.cell(r,3).alignment=wrap
        ws.row_dimensions[r].height=30;r+=1
    r+=1

blk("The tabs",[
 ("1 · Components","THE MAIN SHEET. Every component a client workspace needs, point by point — what it is, why they need it, which surface it lives on, how it opens, what it shows, what they can do, and the rule that governs it."),
 ("2 · Screens x surfaces","The 18 client screens mapped onto Framework B's surfaces (B1-B6), Desktop and Mobile — Joy's format — with the components (CL-…) that appear on each screen."),
 ("3 · Workflows","The 12 jobs a client is trying to get done, the screens that serve each, and the components each job leans on."),
 ("4 · Rules that bind the UI","The non-negotiables from the vision + security docs that any client screen must honour."),
 ("5 · Framework A (alternate)","The same 18 screens under Framework A · Structured Cockpit, kept for comparison. Framework B is the one we are building."),
])
blk("The client in one paragraph",[
 ("Who","An organisation (e.g. Acme Corp) with a few people: requesters who raise work, an approver who authorises spend, and someone in finance. They are outside EZ's boundary."),
 ("What they want","Send work in from wherever they already are, know what it will cost before it runs, be told early if it slips, receive the deliverable into their own store, and be billed only for what they verified."),
 ("What we owe them","Exactly two decisions — approve the scope, approve the delivery. Everything else is pushed to them, not pulled by them."),
 ("What they must never see","Who inside EZ does the work, and any sub-vendor. They see milestones, not the machinery. This is the boundary, and it is a hard rule."),
])
blk("How to read the Components sheet",[
 ("Surface","Which region of the Conversation OS shell it belongs to — B1 nav, B2 list+context, B3 spaces, B4 conversation, B5 prompt, B6 ambient AI, or the right-hand rail."),
 ("Opens as","Inline (always there) · Top drawer (expands in place over the conversation — things about THIS work) · Side drawer (slides from the right rail — things you pull INTO the conversation) · Full screen · Popover."),
 ("Priority","P0 = cannot ship the client experience without it. P1 = needed for a real pilot. P2 = follows once volume justifies it."),
 ("In prototype","Whether it exists in the current client app build."),
])

# ================= COMPONENTS =================
ws=wb.create_sheet("1 · Components")
COLS=[("#",7),("Area",17),("Component",26),("What it is",44),("Why the client needs it",40),
      ("Surface",13),("Opens as",14),("Desktop",34),("Mobile",30),("Key data shown",38),
      ("What they can do",34),("Rule / constraint",40),("P",5),("Built?",9)]
band(ws,len(COLS),"CLIENT WORKSPACE — COMPONENT SPEC   ·   every part that must exist, point by point")
for i,(t,w) in enumerate(COLS,1):
    ws.column_dimensions[get_column_letter(i)].width=w
    c=ws.cell(2,i,t);c.font=hdr;c.fill=PatternFill("solid",fgColor=NAVY);c.border=bd;c.alignment=ctr
ws.row_dimensions[2].height=30
ws.freeze_panes="D3"

C=[
# id, area, component, what, why, surface, opens, desktop, mobile, data, actions, rule, P, built
("CL-01","Shell","Workspace identity","The org's name + mark at the top of the left rail, with a switcher if the person belongs to more than one workspace.","Confirms which company's work they're looking at — people sit in more than one workspace.","B1","Inline","Mark + org name at rail top.","Top bar title.","Org name, logo, role label (Client).","Switch workspace.","One organisation = one workspace. A person can hold membership in many, and they are mutually invisible.","P0","Yes"),
("CL-02","Shell","Section rail","The five domains a client ever needs: Home, Requests, People, Files, Money — plus Help at the foot.","One fixed place to move between the big areas; never more than one click from anywhere.","B1","Inline","Icon rail, labels beneath, badge counts.","Bottom tab bar.","Domain name, unread/awaiting counts.","Jump to a domain.","Money surfaces only appear when the capability is switched on for that tenant.","P0","Yes"),
("CL-03","Shell","Identity & notifications","Avatar and a notification indicator at the foot of the rail.","Know who they're acting as, and that something needs them.","B1","Inline / popover","Avatar + dot at rail foot.","Avatar in top bar.","Name, role, unread count.","Open profile, sign out, open alerts.","Portable root identity — the person authenticates once and acts per membership.","P0","Yes"),
("CL-04","Shell","Global search & command","One field to find any request, invoice, file or person — and to run a command.","Direct access without navigating; the escape hatch for power users.","B1/B5","Popover","Cmd-K opens over the app.","Search control in the header.","Matches grouped by type.","Jump to, or run a command.","Server-side POV scoping — search only ever returns what this client may see.","P1","Partial"),
("CL-05","Shell","Breadcrumb & back","A path back up from wherever they've drilled to.","Never feel lost or trapped two levels deep.","B2","Inline","Back link above the list title.","Native back / drill-stack.","Parent name.","Go up one level.","-","P0","Yes"),

("CL-06","List & context","Request list, grouped","Every request grouped Needs you / Active / Delivered, newest first.","The single place to see the whole book of work.","B2","Inline","Grouped tree in the left pane.","Centre list; tap to open.","Name, status, due, spend, whether it awaits them.","Open a request; filter.","Each node is a room — selecting it re-scopes the conversation.","P0","Yes"),
("CL-07","List & context","Work tree (Request → Assignment → Activity)","The hierarchy of the work itself, rendered with three distinct levels and connector lines.","Understand how a request is actually made up, and talk to any level of it.","B2","Inline","Three indent levels: request (hexagon), assignment (grid + owner chip), activity (status dot).","Breadcrumb + drill list.","Names, status per node, owner initials, mandatory-Deliver tag.","Select any node to talk to it.","Deliver is a mandatory last activity on every assignment. Status rolls up worst-of.","P0","Yes"),
("CL-08","List & context","Transparency dial","A Milestones / Detailed switch controlling how deep the client can see.","Most clients want milestones; some negotiate detail. It must be a setting, not a rebuild.","B2","Inline","Segmented control above the tree.","Segmented control.","Current level.","Switch level (where permitted).","Set per client relationship; the vendor controls the ceiling. Default is milestones.","P0","Yes"),
("CL-09","List & context","Filters & saved views","Chips for All / Awaiting you / Active / At risk / Delivered.","Cut a long list to the few things that matter now.","B2","Inline","Filter chips in the pane.","Filter sheet.","Counts per filter.","Apply, clear, save a view.","-","P1","Yes"),
("CL-10","List & context","New request entry","A primary button that is always visible in the Requests section.","Raising work is the client's most important action — it must never be hunted for.","B2","Inline → full","Primary button under the search field.","Floating action button.","-","Start a request.","Intake also arrives by email, connected tools or API — the button is one of several doors.","P0","Yes"),

("CL-11","Spaces","Attention tiles","A strip of five metrics across the top of the centre, contextual to where they are.","Answer 'what's the state of things' before reading a single message.","B3","Inline","Five tiles; clicking one expands its space.","Two-column tile grid.","Active, awaiting you, delivered, wallet, next due (varies by screen).","Open the underlying space.","Metrics are derived in the same transaction as the state change — no drift, no stale numbers.","P0","Yes"),
("CL-12","Spaces","Plan / milestone timeline","The schedule of a request as milestones on a timeline, with today marked.","Know what's done, what's running, and when it lands.","B3","Top drawer","Expands in place over the conversation; timeline + list toggle.","Owns the centre; timeline becomes a card list, never side-scrolls.","Milestone, state, dates, delivery date, slack.","Switch view; accept a new date.","Border-safe: milestones only. No internal owners, no assignments unless transparency = detailed.","P0","Yes"),
("CL-13","Spaces","Scope & price sheet","Line-by-line what's included, at what level, for what price, by when — plus assumptions.","This is what they're approving. It has to be complete and unambiguous.","B3","Top drawer","Expands over the ground; priced lines and a total.","Full-screen scrollable sheet.","Inclusions, volumes, level, price lines, total, delivery date, assumptions.","Approve, tweak, re-brief.","Nothing starts and nothing is charged before approval. Changing an assumption forces a re-price.","P0","Yes"),
("CL-14","Spaces","Files space","Everything in and out of a request, split Inputs and Outputs.","Supply what the team needs; collect what's been produced.","B3","Top drawer","Expands over the ground; grouped list.","Owns the centre; one card per file.","Filename, direction, status, destination, badges.","Attach, download, revoke access.","Inputs are read on demand from the client's store and never retained. Outputs are written to the client's store as their property.","P0","Yes"),
("CL-15","Spaces","Deliverable space","The delivered artefact, its QA result, what changed, and where it was written.","Everything needed to decide 'do I accept this?' in one view.","B3","Top drawer","Expands over the ground.","Owns the centre; preview + sticky approve.","Files, QA outcome and exceptions, destination, price released on approval.","Download, approve, request revision.","Approving is the client-side confirmation; it is what releases the invoice.","P0","Yes"),
("CL-16","Spaces","Invoice space","One invoice: its lines, each traceable to work the client verified, plus tax and totals.","Trust the bill — see the work behind every line.","B3","Top drawer","Expands over the ground.","Lines as cards.","Lines, dates, tax, total, due date, PO reference.","Pay, query a line, download PDF.","One invoice group = one invoice per entity per month. Rate is snapshotted at invoice time.","P0","Yes"),
("CL-17","Spaces","Wallet ledger","The prepaid balance and every movement against it.","See where the money went and whether work is about to stall.","B3","Top drawer","Expands over the ground; ledger table.","Record-card list.","Lots, holds, settles, refunds, expiry, running balance.","Top up, set auto top-up, export.","Append-only immutable ledger, FIFO across lots: deduct at create, settle at verify, refund on cancel.","P1","Yes"),
("CL-18","Spaces","Service period","For retainers: committed versus used, burn-down, roll-over and close date.","Know if they're under- or over-using what they pay for, before the period closes.","B3","Top drawer","Expands over the ground; burn-down bar.","Owns the centre.","Committed, used, remaining, projected close, roll-over, period value.","Raise work under it; change commitment.","The service period is the billable unit for Type-B — it plays the role Deliver plays for a one-off.","P1","Yes"),
("CL-19","Spaces","Reporting","Spend, volume, turnaround and SLA over a period, with export.","Make the internal case for the relationship.","B3","Top drawer","Charts + table, period filter.","Charts stack; table becomes cards.","Spend by month/offering, request count, avg turnaround, SLA hit-rate, revision rate.","Filter, export, schedule.","Fed by a dedicated reporting store off one lifecycle engine — structurally zero drift.","P1","Yes"),
("CL-20","Spaces","Team & access","The client's own people, their roles and spend limits.","Control who at their org can commit money.","B3","Top drawer","Grouped list with roles.","Card list.","Person, role, spend limit, last active, pending invites.","Add, remove, change role, set a limit.","Roles are server-authoritative and config-driven — never a front-end allow-list.","P1","Yes"),

("CL-21","Conversation","Room stream","The conversation for whatever node is selected — the centre ground of the whole product.","This is where the work is understood, questioned and decided.","B4","Inline (the ground)","Occupies the centre; spaces expand over it.","Detented drawer — peek / half / full.","Messages, decisions, status changes, agent output, in order.","Read, reply, act on cards.","Tree-scoped: each participant sees only the rooms they hold a role in.","P0","Yes"),
("CL-22","Conversation","Decision card","A system message that carries its actions inline — approve, accept a date, request a revision.","The decision happens where the context is; no jumping to a separate screen.","B4/B6","Inline in stream","Bordered card with buttons.","Same card in the drawer.","What happened, what it means, what the options are.","Take the action; dismiss; ask why.","AI suggests, a human approves — the system never acts unilaterally on a commercial decision.","P0","Yes"),
("CL-23","Conversation","Status-change pill","A small chip under a message showing a state transition, e.g. On track → At risk.","See consequence, not just narrative.","B4","Inline in stream","Chip beneath the card body.","Same.","Object name, from-state, to-state.","Click through to the object.","Status is derived worst-of in the same transaction as the change.","P1","Yes"),
("CL-24","Conversation","Human message","A message from a named person, with avatar, role and time.","Talk to actual people, not a ticket system.","B4","Inline in stream","Avatar + bubble.","Same, in the drawer.","Author, role, timestamp, body, references.","Reply, quote, react.","On the client's side of the boundary they see the EZ delivery team, not the individual expert.","P0","Yes"),
("CL-25","Conversation","AI answer (private)","A response only the asker sees, marked 'only you'.","Ask a blunt question without broadcasting it to the room.","B6","Inline in stream","Card marked 'only you'.","Same.","The answer, as a short list, with a follow-on action.","Act on it; discard.","AI's point of view equals the user's point of view — it can never surface what the user may not see.","P1","Yes"),
("CL-26","Conversation","Attachment in stream","A file shared into the room, rendered as a card.","Files belong to the conversation that needed them.","B4","Inline in stream","File card with badges.","Same.","Filename, size, direction, access badge.","Open, download, revoke.","Read-only and revocable for inputs; nothing retained after processing.","P0","Yes"),

("CL-27","Prompt","Message input","The single field that takes a message, a command, or a question.","One place to say anything — no mode-switching.","B5","Inline","Bottom of centre; placeholder adapts to context.","Peek state of the drawer, always visible.","Placeholder tells you what you can do here.","Send.","Context-adaptive: the same field is message, command and query.","P0","Yes"),
("CL-28","Prompt","Scope selector","A control showing exactly what the message will be addressed to — Request, Assignment, Activity or a Person.","Ambiguity about 'who am I telling this to' is the biggest failure mode of a chat-first product.","B5","Popover","Chip left of the input with type + name; opens a picker.","Same chip, sheet picker.","Type label, target name, and for people whether it's a DM.","Change target.","Also determines the boundary — internal vs client-visible — for what you send.","P0","Yes"),
("CL-29","Prompt","Scope indicator line","A quiet line under the input: which room, which boundary, how many people are in it.","Know before sending who will read this.","B5","Inline","Right-aligned under the composer.","Under the peek.","Room path, boundary (internal/with EZ), participant count.","-","Sending into a client room is a boundary crossing and is logged.","P0","Yes"),
("CL-30","Prompt","Attach","Paperclip that brings a file into the room.","Answer a blocker with the asset in one move.","B5","Popover","Icon in the composer.","Icon in the peek.","Recent files, browse, drag-drop.","Attach from device or connected store.","Files are pointers — WS stores a reference, not the bytes.","P0","Yes"),
("CL-31","Prompt","Voice note","Record a spoken message, transcribed into the room.","Faster than typing for a brief or a nuance.","B5","Inline","Mic in the composer.","Mic in the peek.","Duration, transcript.","Record, review, send.","Transcription happens in the sealed enclave; audio is not retained.","P2","Yes"),
("CL-32","Prompt","Slash commands","Typed shortcuts — /status, /brief, /invoice, /export.","Repeat actions without hunting through menus.","B5/B6","Inline + popover","Typing '/' opens a list.","Same.","Command name, what it does, usage count.","Run it.","Commands are saved per person and appear in the Library.","P1","Yes"),
("CL-33","Prompt","Mentions & references","@ to pull in a person, # to reference a file, glossary or request.","Bind a message to the thing it's about.","B5","Popover","Type-ahead on @ and #.","Same.","Matching people and objects.","Insert a reference.","@-mentioning across the boundary notifies on the recipient's own channel.","P1","Partial"),
("CL-34","Prompt","Sticky primary action","On a decision screen the send button becomes the decision — Approve, Pay.","The one thing they're there to do should be the most obvious thing on screen.","B5","Inline","Send button turns into a green Approve.","Sticky action bar above the peek.","Action label.","Take the decision.","Only appears on the two client touchpoints and on payment.","P0","Yes"),

("CL-35","Ambient AI","Proactive risk alert","The system tells the client work is slipping before they ask, with a revised date.","The whole promise is 'you won't need to chase'.","B6","Inline in stream","Alert card with actions.","Same, plus push on their channel.","What slipped, why, the new date, cost impact (usually none).","Accept the date, see what changed, escalate.","SLA monitor raises this automatically; the client is told, not asked to discover it.","P0","Yes"),
("CL-36","Ambient AI","Price explanation & alternates","When they push back on price, it explains the lines and offers cheaper or faster shapes.","Turns a negotiation into a choice.","B6","Inline in stream","Answer card with options.","Same.","Each alternate with its delta and trade-off.","Swap to an alternate.","Suggest, never act — a re-price still needs their approval.","P1","Yes"),
("CL-37","Ambient AI","Status roll-up on demand","Ask 'where is everything' and get a cross-request answer.","One question instead of opening five rooms.","B6","Inline in stream","Private answer card.","Same.","Per-request state, what needs them, what's at risk.","Jump to any item.","Border-safe — rolls up milestones, never internal steps.","P1","Yes"),
("CL-38","Ambient AI","Revision vs new scope","Classifies a change request as a free revision or new priced scope, before they send it.","Prevents the most common commercial argument.","B6","Inline in stream","Explanation card.","Same.","Which it is and why.","Send as revision, or ask for a quote.","Goodwill revision creates a new linked assignment with no claw-back.","P1","Yes"),
("CL-39","Ambient AI","Wallet runway warning","Warns that the balance will run out before committed work settles.","Work stalling for a funding reason is an avoidable failure.","B6","Inline in stream","Alert card with top-up action.","Same.","Burn rate, days of runway, what would stall.","Top up, set auto top-up.","Low-balance alerts fire before work stalls, not at the point of failure.","P1","Yes"),

("CL-40","Right rail","Library drawer","Saved commands, request templates and playbooks.","Repeat what worked last time without rebuilding the brief.","Right rail","Side drawer","Icon rail on the right; 284px drawer.","Full-height sheet.","Command name, template contents, usage count.","Insert into composer, reuse a template.","Templates are config, versioned and safe to rename.","P1","Yes"),
("CL-41","Right rail","Files drawer","References applied automatically plus recent files in this room.","Grab a file while still typing.","Right rail","Side drawer","Drawer beside the conversation.","Sheet.","Filename, what it is, applied badge.","Insert, open, manage access.","Same brokered-access rule as CL-14.","P1","Yes"),
("CL-42","Right rail","People drawer","Who is in this room — their side and EZ's side, with presence.","Know your audience before you write.","Right rail","Side drawer","Drawer with two groups.","Sheet.","Name, role, presence, side of the boundary.","DM, add someone from their org.","Shows the EZ team, never individual experts or sub-vendors.","P0","Yes"),
("CL-43","Right rail","Alerts drawer","Everything that has pinged them, colour-coded by severity.","One inbox for attention, separate from the room they're in.","Right rail","Side drawer","Drawer, badge on the icon.","Sheet + push.","Event, object, time, severity.","Open the source, mark read.","Delivery is event-driven with retry and audit; both channel switches must be on.","P1","Yes"),
("CL-44","Right rail","Activity log drawer","Every change to this object, in order, attributed.","Answer 'who changed what, when' without asking.","Right rail","Side drawer","Vertical timeline.","Sheet.","Actor, action, timestamp.","Filter, export.","Tenant-inspectable and tamper-evident — the client can verify us.","P1","Yes"),

("CL-45","Money","Invoice list (the invert)","A table of every invoice with selection and bulk actions.","Finance works in lists, not conversations.","B4 (inverted)","Inline (table takes the ground)","Table owns the centre; conversation collapses to a tile.","Record-card list, never side-scrolling.","Number, period, amount, due, status, age.","Select, pay, export, remind.","On transactional surfaces the ground inverts: table takes the centre, chat becomes a tile.","P0","Yes"),
("CL-46","Money","Pay","Settle an invoice from wallet, card or transfer.","The point of the whole finance area.","B5","Inline / full","Sticky Pay in the composer row.","Sticky action.","Amount, method, date.","Pay, schedule, part-pay.","Invoice and payout are decoupled — the expert is paid on verified work regardless.","P0","Partial"),
("CL-47","Money","Query a line","Dispute or question a single invoice line, in a thread attached to that line.","Resolve a billing question without stopping the whole invoice.","B4","Inline in stream","Query opens a thread on the line.","Drawer.","The line, the verified work behind it, the thread.","Raise a query, accept the explanation.","A dispute never blocks the performer payout for verified work.","P1","Partial"),
("CL-48","Money","Top up wallet","Add prepaid credit, optionally on a rule.","Keep work flowing.","B3/B5","Top drawer + action","Action in the wallet space.","Sheet.","Amount, method, resulting balance and runway.","Top up, set auto top-up.","Credits are client-side; coins are the expert-side normaliser. They are separate scales.","P1","Partial"),

("CL-49","Lifecycle","Intake / compose","The brief: what they need, files, deadline — as a message, not a form wall.","Describing the need must be as easy as writing an email.","B4/B5","Inline in stream","Compose card in the stream; AI asks only what it can't infer.","Full-screen stepped compose.","Ask, files, deadline, inferred offering and level, live estimate.","Send to scope, save draft.","Intake also arrives via email (1 thread = 1 request), connected tools or API. Nothing is ever dropped.","P0","Yes"),
("CL-50","Lifecycle","Scope approval (touchpoint 1)","The client's authorisation of scope, price and date.","Human touchpoint 1 of 2. Nothing runs without it.","B3+B4+B5","Top drawer + sticky action","Scope sheet above, decision card in the stream, Approve as the send action.","Sheet + sticky Approve.","Everything in CL-13.","Approve, tweak, re-brief.","No auto-start. AI proposes; a human ratifies.","P0","Yes"),
("CL-51","Lifecycle","Clarification / unblock","A question from the team that is blocking work, with the answer inline.","A blocked step is the most expensive state; clearing it must take seconds.","B4","Inline in stream","Question quoted, reply and attach inline.","Drawer opens at half on the question.","What's blocked, what's needed, impact on the date.","Reply, attach, say they don't have it.","Escalation ladder if unanswered: in-app, then email, then WhatsApp/SMS.","P0","Yes"),
("CL-52","Lifecycle","Delivery approval (touchpoint 2)","The client accepts the delivered work.","Human touchpoint 2 of 2 — and the commercial trigger.","B3+B4+B5","Top drawer + sticky action","Deliverable space above, decision card in the stream.","Sheet + sticky Approve.","Everything in CL-15, plus what approving will release.","Approve, request revision, download first.","Verify is an explicit human gate before billing. No auto-invoice on 'marked complete'.","P0","Yes"),
("CL-53","Lifecycle","Revision request","Send work back with a note; opens a linked revision.","Fairness when the work isn't right.","B4","Inline in stream","Note becomes the revision brief; system confirms scope and date.","Drawer at full to write it.","What's wrong, what's in scope, charge (none), new date.","Send, add detail, escalate.","Goodwill revision = new linked assignment, no charge, no claw-back. Original delivery stays valid.","P0","Yes"),
("CL-54","Lifecycle","Engagement container","The thing above a request — one-off, retainer, or outcome.","Retainer clients don't think in single requests.","B2/B3","Inline + top drawer","Engagement node with a child per period.","List of periods.","Delivery model, billing model, cadence, periods and their state.","Raise work under it, review a period.","A new commercial shape is a new config combination, never new code.","P1","Yes"),

("CL-55","Governance","Roles & permissions","Requester, approver, finance — who can do what.","Stop anyone raising unlimited spend.","B3","Top drawer","People space with role column.","Card list.","Person, role, what it permits.","Assign, change, revoke.","Server-authoritative, config-driven RBAC.","P1","Yes"),
("CL-56","Governance","Spend limits","A cap per person or per request.","Real financial control, not a policy document.","B3","Top drawer","Limit field per person.","Sheet.","Limit, used, remaining.","Set, change, remove.","Enforced at approval time, not after the fact.","P1","Partial"),
("CL-57","Governance","Delivery destination","Where deliverables are written — their SharePoint, Drive, or own infrastructure.","Deliverables must land in their world, as their property.","B3","Top drawer","Setting with connection state.","Sheet.","Connected store, path, test result.","Connect, change, test.","Outputs are transferred to the client's store; WS keeps a pointer and the attribution only.","P0","Partial"),
("CL-58","Governance","Notification policy","Which events reach them, on which channel.","Stop the noise without missing the two that matter.","B3","Top drawer","Event x channel matrix.","Grouped list.","Event, channels on/off, recipients.","Toggle per event and channel.","Recipient rules are config by role and status — never a hard-coded domain allow-list.","P1","Yes"),
("CL-59","Governance","Glossary & references","Term bases, style guides and brand assets applied automatically to their work.","Consistency without repeating themselves on every brief.","B3/Right rail","Top drawer + side drawer","Files space and Library drawer.","Sheet.","File, what it is, whether it's applied.","Upload, apply, retire.","Held as per-client reference data for in-scope AI use only.","P1","Yes"),
("CL-60","Governance","Transparency & boundary setting","How much of the work graph this client is allowed to see.","Different clients negotiate different visibility.","B3","Top drawer","Setting with a live preview.","Sheet.","Current level, what it reveals.","Change (within the ceiling EZ sets).","Default opaque/milestones. The vendor controls the ceiling; the client cannot raise it themselves.","P1","Partial"),

("CL-61","Trust","Brokered input indicator","A visible 'read-only, revocable' badge on every file they've supplied.","Prove we don't hold their data.","B3/B4","Inline","Badge on the file card + explanatory note.","Same.","Access state, who can read, revoke control.","Revoke access.","Read on demand from the sender, revocable, nothing retained. This is a security guarantee, not copy.","P0","Yes"),
("CL-62","Trust","Deliverable ownership statement","A line stating the output was written to their store as their property.","The commercial and legal relationship made visible.","B3","Inline","Note under the deliverable.","Same.","Destination, ownership, what WS keeps.","Open in their store.","WS keeps the pointer and the attribution; the client holds the asset.","P0","Yes"),
("CL-63","Trust","Boundary indicator","A quiet explanation of why they see milestones and not people.","Pre-empts 'why can't I see who's working on this'.","B2/B3","Inline","Note under the plan and in the People drawer.","Same.","What they see and what they don't.","-","Full encapsulation both directions; chained vendors are invisible.","P1","Yes"),
("CL-64","Trust","Audit visibility","The client can inspect the log of actions on their own data.","'Don't trust us, verify' — a differentiator, not a checkbox.","Right rail","Side drawer","Activity log drawer.","Sheet.","Actor, action, time, object.","Filter, export.","Tamper-evident and tenant-inspectable, including operator actions.","P1","Partial"),

("CL-65","Onboarding","Readiness checklist","What still needs setting before they can transact.","A half-configured workspace fails silently at the worst moment.","B3","Top drawer","Checklist tiles with gaps flagged.","Checklist grid.","Each step, its state, what it unlocks.","Complete a step.","Mirrors the internal readiness engine — capability gates are computed, not asserted.","P1","Yes"),
("CL-66","Onboarding","First-run guidance","A welcome thread explaining the flow and the two decisions they own.","Understand the model in under a minute, without training.","B4","Inline in stream","Welcome thread with actions.","Drawer at half.","How work flows, what's asked of them.","Do a setup step, take the tour.","Ambient and inline — never a modal tour or a chatbot panel.","P1","Yes"),
("CL-67","Onboarding","Starter templates","Three example requests they can adapt.","The blank page is the enemy of first use.","B2/Right rail","Inline + side drawer","Starter cards in the empty state; templates in the Library.","Cards.","Template name, what it contains.","Use it.","Seeded from an industry archetype — generic knowledge, never another client's setup.","P2","Yes"),

("CL-68","Mobile","Detented conversation drawer","On mobile the ground flips: spaces own the screen and the conversation becomes a drawer.","Small screens can't do side-by-side; one must dominate.","B4","Side/bottom drawer","n/a","Peek / half / full, dragged or tapped.","Conversation, with the prompt always visible at peek.","Expand, collapse, send from peek.","The prompt is always visible as the peek state — you can always act.","P0","Yes"),
("CL-69","Mobile","Record-card lists","Wide tables become one card per row.","Side-scrolling a table on a phone is a broken experience.","B3/B4","Inline","n/a","Every table renders as cards.","The same fields, stacked.","Same actions per card.","Never side-scroll a table on mobile — this is a hard rule.","P0","Yes"),
("CL-70","Mobile","Push to their channel","Approvals and risk alerts reach them on WhatsApp, SMS or email.","Clients don't sit in our app all day.","B6","External","n/a","Push notification deep-links into the exact decision.","What needs them, one-tap to act.","Approve or open.","Channel-native by config; WS-chat default, WhatsApp/SMS as paid fallbacks.","P1","Partial"),
]

r=3
AREACOL={"Shell":GREY,"List & context":SKYSOFT,"Spaces":SOFT,"Conversation":"FFF3E8","Prompt":SKYSOFT,
 "Ambient AI":"EAF7FE","Right rail":GREY,"Money":GOODS,"Lifecycle":SOFT,"Governance":GREY,
 "Trust":WARNS,"Onboarding":SKYSOFT,"Mobile":GREY}
for row in C:
    for i,v in enumerate(row,1):
        c=ws.cell(r,i,v);c.border=bd;c.alignment=wrap;c.font=body
        if i==1:c.font=Font(name=F,size=8.5,bold=True,color=NAVY);c.alignment=ctr;c.fill=PatternFill("solid",fgColor=AREACOL.get(row[1],GREY))
        if i==2:c.font=sub;c.fill=PatternFill("solid",fgColor=AREACOL.get(row[1],GREY));c.alignment=Alignment(wrap_text=True,vertical="center")
        if i==3:c.font=bodyb
        if i in (6,7,13,14):c.alignment=ctr;c.font=Font(name=F,size=8.5,bold=(i in(13,14)),color=NAVY)
        if i==12:c.font=mut
        if i==13:
            c.fill=PatternFill("solid",fgColor={"P0":CRITS,"P1":WARNS,"P2":GREY}.get(v,GREY))
        if i==14:
            c.fill=PatternFill("solid",fgColor={"Yes":GOODS,"Partial":WARNS,"No":CRITS}.get(v,GREY))
    ws.row_dimensions[r].height=62
    r+=1

# ================= SCREENS x SURFACES (carried from Joy's mapping) =================
from openpyxl import load_workbook
SRC=load_workbook("/tmp/claude-0/-home-user--WS-3-0/6fa8f15a-8d85-5165-9628-51eda7630d35/scratchpad/WS3_Client_Screen_Surface_Mapping.xlsx")

# which components appear on each screen
ONSCREEN={
"C-01":"CL-01,02,03,05 · 11,65,66,67 · 21,27,28,29 · 57,58 · 62",
"C-02":"CL-01,02,03,04 · 06,09,11 · 21,22,23,35 · 27,28,29,34 · 43,40 · 68",
"C-03":"CL-10,49 · 11 · 24,26 · 27,28,30,31,33 · 36 · 14,61 · 67",
"C-04":"CL-07,08 · 12,13 · 22,25,36 · 27,28,29,34 · 50 · 56,63 · 68",
"C-05":"CL-07,08,12 · 11,14 · 21,23,24,25 · 27,28,29,30,31,32,33 · 35,37 · 40,41,42,43,44 · 63",
"C-06":"CL-07,12 · 21,23 · 27,28,29 · 35,37 · 44 · 69",
"C-07":"CL-06,07 · 14 · 21,24,26 · 27,28,30,33 · 25 · 41,42 · 51,61",
"C-08":"CL-14,15 · 21,23,26 · 27,28,30,33 · 40,41 · 57,59,61,62",
"C-09":"CL-07 · 15 · 22,25 · 27,28,34 · 52 · 62,64 · 68",
"C-10":"CL-07 · 15 · 21,24 · 27,28 · 38 · 53",
"C-11":"CL-06,09,11 · 21 · 27,28,29 · 37 · 43 · 69",
"C-12":"CL-45,46 · 11 · 21,27,28 · 39 · 44 · 69",
"C-13":"CL-16 · 45,46,47 · 21,24 · 27,28,34 · 52,64",
"C-14":"CL-17,48 · 21,23 · 27,28 · 39 · 44 · 69",
"C-15":"CL-18,54 · 07,12 · 21,23 · 27,28 · 37 · 42",
"C-16":"CL-19 · 09 · 21 · 27,28 · 37 · 69",
"C-17":"CL-20,55,56 · 21,23 · 27,28 · 42,44 · 63",
"C-18":"CL-57,58,59,60 · 21,23 · 27,28 · 40 · 62,63,70",
}
NOTE={
"C-01":"Nothing here asks for a decision — it asks for readiness.",
"C-02":"'Needs you' is the only ranked thing; everything else is ambient.",
"C-03":"Composing IS a message. There is no form.",
"C-04":"★ Human touchpoint 1. Nothing runs and nothing is charged before this.",
"C-05":"The default state of the product. Conversation holds the ground.",
"C-06":"Plan expands in place over the room — it never navigates away.",
"C-07":"Where a blocker gets cleared in one exchange rather than an email chain.",
"C-08":"Inputs are brokered; outputs land in the client's own store.",
"C-09":"★ Human touchpoint 2 — and the commercial trigger. Approving releases the invoice.",
"C-10":"Revision is goodwill and free; new scope is re-priced and re-approved.",
"C-11":"List view. Conversation stays as a digest, not a per-request stream.",
"C-12":"THE INVERT — the table takes the ground, conversation collapses to a tile.",
"C-13":"Every line traces back to work the client themselves verified.",
"C-14":"Append-only ledger: deduct at create, settle at verify, refund on cancel.",
"C-15":"For retainers. The engagement sits above individual requests.",
"C-16":"The client's case to their own boss. Read-only.",
"C-17":"Access control that actually binds — limits are enforced at approval time.",
"C-18":"Config, not code — every switch here is declarative and versioned.",
}

src=SRC["Framework B"]
ws=wb.create_sheet("2 · Screens x surfaces")
band(ws,16,"FRAMEWORK B · CONVERSATION OS   ·   the 18 client screens across the six surfaces   ·   ★ = a client decision")
W=[7,26]+[30]*12+[34,40]
for i,w in enumerate(W,1): ws.column_dimensions[get_column_letter(i)].width=w
SURF=["B1 · Section nav","B2 · List + context","B3 · Spaces","B4 · Conversation","B5 · Command prompt","B6 · Ambient AI"]
SCOL={"B1 · Section nav":GREY,"B2 · List + context":SKYSOFT,"B3 · Spaces":SOFT,
      "B4 · Conversation":"FFF3E8","B5 · Command prompt":SKYSOFT,"B6 · Ambient AI":"EAF7FE"}
for j,(t,) in enumerate([(x,) for x in ["#","Screen"]],1):
    ws.merge_cells(start_row=2,start_column=j,end_row=3,end_column=j)
    c=ws.cell(2,j,t);c.font=hdr;c.fill=PatternFill("solid",fgColor=NAVY);c.alignment=ctr
    for rr in (2,3): ws.cell(rr,j).border=bd
for k,s in enumerate(SURF):
    c1=3+k*2
    ws.merge_cells(start_row=2,start_column=c1,end_row=2,end_column=c1+1)
    c=ws.cell(2,c1,s);c.font=Font(name=F,size=9,bold=True,color=NAVY)
    c.fill=PatternFill("solid",fgColor=SCOL[s]);c.alignment=ctr
    for d,lab in enumerate(["Desktop","Mobile"]):
        c=ws.cell(3,c1+d,lab);c.font=sub;c.fill=PatternFill("solid",fgColor=SCOL[s]);c.alignment=ctr;c.border=bd
    ws.cell(2,c1).border=bd;ws.cell(2,c1+1).border=bd
for j,t in [(15,"Components on this screen"),(16,"Why this screen exists")]:
    ws.merge_cells(start_row=2,start_column=j,end_row=3,end_column=j)
    c=ws.cell(2,j,t);c.font=hdr;c.fill=PatternFill("solid",fgColor=ORANGE);c.alignment=ctr
    for rr in (2,3): ws.cell(rr,j).border=bd
ws.row_dimensions[2].height=24;ws.row_dimensions[3].height=18
r=4
for row in src.iter_rows(min_row=4,values_only=True):
    if not row[0]: continue
    sid=row[0]
    for i,v in enumerate(row,1):
        c=ws.cell(r,i,v if v else "—");c.border=bd;c.alignment=wrap;c.font=body
        if i==1:c.font=Font(name=F,size=9,bold=True,color=NAVY);c.alignment=ctr;c.fill=PatternFill("solid",fgColor=GREY)
        if i==2:c.font=bodyb;c.fill=PatternFill("solid",fgColor=SOFT if "★" in str(v) else GREY)
        if i>=3:
            c.fill=PatternFill("solid",fgColor="FFFFFF" if (i-3)//2%2 else "FCFDFE")
            if i in (9,10):c.font=Font(name=F,size=9,color="1A1A1A",bold=False)
    c=ws.cell(r,15,ONSCREEN.get(sid,"—"));c.border=bd;c.alignment=wrap
    c.font=Font(name=F,size=8.5,bold=True,color="C05F17");c.fill=PatternFill("solid",fgColor=SOFT)
    c=ws.cell(r,16,NOTE.get(sid,""));c.border=bd;c.alignment=wrap;c.font=mut
    ws.row_dimensions[r].height=74;r+=1
ws.freeze_panes="C4"

# ================= WORKFLOWS =================
src=SRC["Client workflows"]
WFC={
"Get set up":"CL-01,02,03 · 57,58,65,66,67 · 20,55",
"Send work in":"CL-10,49 · 27,28,30,31,33 · 14,26,61 · 36",
"Know what it will cost before it runs ★":"CL-12,13 · 22,25,36 · 34,50 · 56",
"Know it's on track without chasing":"CL-07,08,12 · 21,23,35,37 · 43,63",
"Answer questions fast":"CL-21,24,26 · 27,30,33 · 41,42 · 51 · 70",
"Receive and check the work ★":"CL-15 · 22,34,52 · 57,62,64",
"Get it fixed if it's not right":"CL-38,53 · 15,21,24",
"Understand what I'm paying":"CL-45,46,47 · 16 · 44,64",
"Keep credit topped up":"CL-17,39,48 · 23",
"Run an ongoing engagement":"CL-18,54 · 07,12,37",
"Prove the value internally":"CL-19 · 09,37",
"Control who can spend":"CL-20,55,56 · 44,63",
}
ws=wb.create_sheet("3 · Workflows")
band(ws,7,"CLIENT WORKFLOWS   ·   the 12 jobs they are trying to get done   ·   ★ = a client decision")
COLS3=[("Workflow (job to be done)",30),("What the client wants",44),("Trigger",22),
       ("Steps",56),("Screens",16),("Components used",34),("Outcome",26)]
for i,(t,w) in enumerate(COLS3,1):
    ws.column_dimensions[get_column_letter(i)].width=w
    c=ws.cell(2,i,t);c.font=hdr;c.fill=PatternFill("solid",fgColor=NAVY);c.border=bd;c.alignment=ctr
ws.row_dimensions[2].height=26
r=3
for row in src.iter_rows(min_row=3,values_only=True):
    if not row[0]: continue
    vals=list(row[:5])+[WFC.get(row[0],"—"),row[5]]
    for i,v in enumerate(vals,1):
        c=ws.cell(r,i,v if v else "—");c.border=bd;c.alignment=wrap;c.font=body
        if i==1:c.font=bodyb;c.fill=PatternFill("solid",fgColor=SOFT if "★" in str(v) else GREY)
        if i==3:c.font=mut
        if i==5:c.font=Font(name=F,size=8.5,bold=True,color=NAVY);c.alignment=ctr;c.fill=PatternFill("solid",fgColor=SKYSOFT)
        if i==6:c.font=Font(name=F,size=8.5,bold=True,color="C05F17");c.fill=PatternFill("solid",fgColor=SOFT)
        if i==7:c.font=Font(name=F,size=9,bold=True,color="0A351F");c.fill=PatternFill("solid",fgColor=GOODS)
    ws.row_dimensions[r].height=66;r+=1
ws.freeze_panes="B3"

# ================= RULES =================
ws=wb.create_sheet("4 · Rules that bind the UI")
COLS2=[("#",6),("Rule",42),("What it means for a client screen",70),("Where it shows up",34)]
band(ws,4,"NON-NEGOTIABLE RULES — any client screen must honour these",color=NAVY)
for i,(t,w) in enumerate(COLS2,1):
    ws.column_dimensions[get_column_letter(i)].width=w
    c=ws.cell(2,i,t);c.font=hdr;c.fill=PatternFill("solid",fgColor=NAVY);c.border=bd;c.alignment=ctr
RULES=[
("R-01","Exactly two human touchpoints","The client is asked to decide twice per request: approve the scope, approve the delivery. Every other state change is pushed to them. If a screen asks for a third decision, it is wrong.","CL-50, CL-52"),
("R-02","Border-safe by default","The client sees milestones. Never internal assignments, expert names, or sub-vendors — unless the relationship's transparency dial is explicitly raised.","CL-07, CL-08, CL-12, CL-42, CL-63"),
("R-03","Nothing runs, nothing is charged, before approval","No work starts and no money moves until the scope is approved. Price is fixed at approval and snapshotted.","CL-13, CL-50"),
("R-04","Verify is the commercial trigger","Approving the deliverable is what releases the invoice. There is no auto-invoice on 'marked complete'.","CL-15, CL-52, CL-45"),
("R-05","Inputs are brokered, never retained","Client files are read on demand from their store, access is scoped and revocable, and nothing is kept after processing.","CL-14, CL-26, CL-61"),
("R-06","Deliverables are the client's property","Outputs are written to the client's own store. WS keeps a pointer and the attribution only.","CL-15, CL-57, CL-62"),
("R-07","Deliver is mandatory-last","Every assignment ends in a Deliver activity. Status rolls up worst-of, in the same transaction as the change.","CL-07, CL-12"),
("R-08","Suggest, never act","AI composes, prices, proposes and warns. A human always ratifies anything commercial.","CL-22, CL-25, CL-36, CL-50"),
("R-09","AI's point of view = the user's","The assistant can never surface anything the person is not entitled to see.","CL-25, CL-37"),
("R-10","Revision is free, new scope is priced","A goodwill revision creates a linked assignment with no charge and no claw-back; new scope is re-priced and re-approved.","CL-38, CL-53"),
("R-11","Money is decoupled","The client invoice and the expert payout are two independent sides of one billable event. A client dispute never blocks a payout for verified work.","CL-45, CL-46, CL-47"),
("R-12","Wallet is an append-only ledger","Deduct at create, settle at verify, refund on cancel, FIFO across lots. The ledger is the single source of truth.","CL-17, CL-48"),
("R-13","Capability gating","If a capability is off for that tenant, its fields, screens and endpoints do not exist — they are not hidden, they are absent.","CL-02, CL-17, CL-19"),
("R-14","Config, not code","Terminology, roles, notification rules, transparency and templates are all declarative config, versioned and effective-dated.","CL-08, CL-55, CL-58, CL-60"),
("R-15","Mobile flips the ground","On mobile, spaces own the centre and conversation becomes a detented drawer. The prompt is always visible as the peek state.","CL-68"),
("R-16","Never side-scroll a table on mobile","Wide tables render as one card per row.","CL-69"),
("R-17","Ambient AI, never a chatbot panel","Assistance lives inline in the stream and in the prompt. There is no separate assistant sidebar.","CL-22, CL-25, CL-35"),
("R-18","Tenant-inspectable audit","The client can verify what happened on their own data, including actions taken by the operator.","CL-44, CL-64"),
]
r=3
for x in RULES:
    for i,v in enumerate(x,1):
        c=ws.cell(r,i,v);c.border=bd;c.alignment=wrap;c.font=body
        if i==1:c.font=Font(name=F,size=8.5,bold=True,color=NAVY);c.alignment=ctr;c.fill=PatternFill("solid",fgColor=WARNS)
        if i==2:c.font=bodyb;c.fill=PatternFill("solid",fgColor=SOFT)
        if i==4:c.font=Font(name=F,size=8.5,bold=True,color="901918");c.alignment=ctr
    ws.row_dimensions[r].height=48;r+=1
ws.freeze_panes="A3"

# ================= FRAMEWORK A (alternate, for comparison) =================
src=SRC["Framework A"]
ws=wb.create_sheet("5 · Framework A (alternate)")
band(ws,16,"FRAMEWORK A · STRUCTURED COCKPIT   ·   the same 18 screens, the other way   ·   kept for comparison only",color=NAVY)
WA=[7,26]+[30]*14
for i,w in enumerate(WA,1): ws.column_dimensions[get_column_letter(i)].width=w
SURFA=["A1 · Section nav","A2 · Attention bar","A3 · Drill navigation","A4 · Focus surface",
       "A5 · Context facets","A6 · Command & search","A7 · Ambient AI"]
for j,t in [(1,"#"),(2,"Screen")]:
    ws.merge_cells(start_row=2,start_column=j,end_row=3,end_column=j)
    c=ws.cell(2,j,t);c.font=hdr;c.fill=PatternFill("solid",fgColor=NAVY);c.alignment=ctr
    for rr in (2,3): ws.cell(rr,j).border=bd
for k,s in enumerate(SURFA):
    c1=3+k*2
    ws.merge_cells(start_row=2,start_column=c1,end_row=2,end_column=c1+1)
    c=ws.cell(2,c1,s);c.font=Font(name=F,size=9,bold=True,color=NAVY)
    c.fill=PatternFill("solid",fgColor=GREY if k%2 else SKYSOFT);c.alignment=ctr
    for d,lab in enumerate(["Desktop","Mobile"]):
        c=ws.cell(3,c1+d,lab);c.font=sub;c.fill=PatternFill("solid",fgColor=GREY if k%2 else SKYSOFT);c.alignment=ctr;c.border=bd
    ws.cell(2,c1).border=bd;ws.cell(2,c1+1).border=bd
ws.row_dimensions[2].height=24;ws.row_dimensions[3].height=18
r=4
for row in src.iter_rows(min_row=4,values_only=True):
    if not row[0]: continue
    for i,v in enumerate(row,1):
        c=ws.cell(r,i,v if v else "—");c.border=bd;c.alignment=wrap;c.font=body
        if i==1:c.font=Font(name=F,size=9,bold=True,color=NAVY);c.alignment=ctr;c.fill=PatternFill("solid",fgColor=GREY)
        if i==2:c.font=bodyb;c.fill=PatternFill("solid",fgColor=SOFT if "★" in str(v) else GREY)
    ws.row_dimensions[r].height=62;r+=1
ws.freeze_panes="C4"

OUT="/home/user/-WS-3.0/WS3_Client_Workspace_Spec.xlsx"
wb.save(OUT)
print("saved", OUT, "·", len(C), "components ·", len(RULES), "rules · tabs:", wb.sheetnames)
